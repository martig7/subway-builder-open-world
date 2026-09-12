#!/usr/bin/env python3
"""Extract a non-additive foreign-visitor comparison for metro PT pairs."""

from __future__ import annotations

import argparse
import calendar
import csv
import hashlib
import json
from decimal import Decimal
from pathlib import Path


YEAR = 2024
SOURCE_FILENAME = "ffdata-2024-nationality-od.xlsx"
SOURCE_URL = "https://www.mlit.go.jp/sogoseisaku/soukou/content/2024_Nationality_OD.xlsx"
SOURCE_SHA256 = "05b5b9b9bfb01b91c44a99be4991d8bfb00d364e397636b1c7e307655e1a0a2a"
SHEET_NAME = "全国籍全機関"
REGIONS = {
    "tokyo": frozenset({"11", "12", "13", "14"}),
    "chukyo": frozenset({"21", "23", "24"}),
    "kinki": frozenset({"26", "27", "28", "29"}),
}
CSV_FIELDS = [
    "region",
    "origin_prefecture_code",
    "origin_prefecture_name",
    "destination_prefecture_code",
    "destination_prefecture_name",
    "foreign_visitor_legs_per_year",
    "calendar_day_average_visitor_legs",
    "year",
    "source",
    "scope",
    "comparison_status",
]


def repository_root() -> Path:
    return Path(__file__).resolve().parents[4]


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def prefecture_code(value: object) -> str:
    if not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError(f"invalid FF-Data prefecture code: {value!r}")
    code = int(value)
    if not 2 <= code <= 47:
        raise ValueError(f"unexpected FF-Data prefecture code in metro pair: {code}")
    return f"{code:02d}"


def annual_visitor_legs(value: object) -> int:
    if not isinstance(value, (int, float)) or value < 0:
        raise ValueError(f"missing or invalid FF-Data cell: {value!r}")
    persons = Decimal(str(value)) * Decimal(1000)
    if persons != persons.to_integral_value():
        raise ValueError(f"FF-Data thousand-person cell is not integral persons: {value!r}")
    return int(persons)


def extract_rows(workbook_path: Path, prefecture_names: dict[str, str]) -> list[dict[str, object]]:
    import openpyxl

    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = book[SHEET_NAME]
        if sheet["B3"].value != f"{YEAR}年" or sheet["B7"].value != "（千人/年）":
            raise ValueError("FF-Data year or unit header changed")
        destination_columns: dict[str, int] = {}
        for column in range(6, 58):
            raw_code = sheet.cell(9, column).value
            if isinstance(raw_code, (int, float)) and 2 <= int(raw_code) <= 47:
                destination_columns[prefecture_code(raw_code)] = column
        origin_rows: dict[str, int] = {}
        for row in range(12, 64):
            raw_code = sheet.cell(row, 3).value
            if isinstance(raw_code, (int, float)) and 2 <= int(raw_code) <= 47:
                origin_rows[prefecture_code(raw_code)] = row
        if len(destination_columns) != 46 or len(origin_rows) != 46:
            raise ValueError("FF-Data prefecture axes are incomplete or duplicated")

        days = 366 if calendar.isleap(YEAR) else 365
        rows: list[dict[str, object]] = []
        for region, members in REGIONS.items():
            for origin in sorted(members):
                for destination in sorted(members - {origin}):
                    source_value = sheet.cell(
                        origin_rows[origin], destination_columns[destination]
                    ).value
                    annual = annual_visitor_legs(source_value)
                    average = Decimal(annual) / Decimal(days)
                    rows.append(
                        {
                            "region": region,
                            "origin_prefecture_code": origin,
                            "origin_prefecture_name": prefecture_names[origin],
                            "destination_prefecture_code": destination,
                            "destination_prefecture_name": prefecture_names[destination],
                            "foreign_visitor_legs_per_year": annual,
                            "calendar_day_average_visitor_legs": format(
                                average.quantize(Decimal("0.000001")), "f"
                            ),
                            "year": YEAR,
                            "source": "MLIT FF-Data 2024",
                            "scope": "inbound foreign visitors; domestic-visit-to-domestic-visit inter-prefecture legs",
                            "comparison_status": "not_combined_with_resident_pt_weekday_baseline",
                        }
                    )
    finally:
        book.close()
    if len(rows) != 30:
        raise ValueError(f"expected 30 directed metro pairs, found {len(rows)}")
    return rows


def main() -> int:
    root = repository_root()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=root / "map-creator/data/sources/japan/noncommute/raw" / SOURCE_FILENAME,
    )
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    output_dir = args.output_dir.resolve()
    actual_hash = hash_file(workbook_path)
    if actual_hash != SOURCE_SHA256:
        raise ValueError(f"FF-Data source SHA-256 mismatch: {actual_hash}")
    prefecture_names = json.loads(
        (root / "worlds/japan/geography/prefecture-display-names.json").read_text(
            encoding="utf-8"
        )
    )
    rows = extract_rows(workbook_path, prefecture_names)
    output_dir.mkdir(parents=True, exist_ok=True)
    with (output_dir / "foreign-visitor-metro-2024.csv").open(
        "w", encoding="utf-8", newline=""
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)

    annual_by_region = {
        region: sum(
            int(row["foreign_visitor_legs_per_year"])
            for row in rows
            if row["region"] == region
        )
        for region in REGIONS
    }
    days = 366 if calendar.isleap(YEAR) else 365
    report = {
        "schemaVersion": 1,
        "source": {
            "url": SOURCE_URL,
            "filename": SOURCE_FILENAME,
            "sha256": actual_hash,
            "bytes": workbook_path.stat().st_size,
            "sheet": SHEET_NAME,
            "sourceUnit": "thousand persons per year",
        },
        "year": YEAR,
        "calendarDays": days,
        "status": "separate_non_additive_comparison; not in movements.csv or map-data.json",
        "directedPairCount": len(rows),
        "annualVisitorLegsByRegion": annual_by_region,
        "annualVisitorLegs": sum(annual_by_region.values()),
        "calendarDayAverageVisitorLegs": format(
            (Decimal(sum(annual_by_region.values())) / Decimal(days)).quantize(
                Decimal("0.000001")
            ),
            "f",
        ),
    }
    (output_dir / "foreign-visitor-metro-2024-report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
