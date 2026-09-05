#!/usr/bin/env python3
"""Materialize official e-Stat demand evidence one Japanese prefecture at a time.

This compiles statistical inputs only. It deliberately labels clustered mesh
locations as demand-site candidates rather than buildings; footprint allocation
happens later in the map compiler.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import statistics
import time
import zipfile
from collections import defaultdict
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Iterable

from pyproj import Transformer
from shapely import make_valid
from shapely.geometry import Point, box, mapping, shape
from shapely.ops import nearest_points, transform, unary_union
from shapely.strtree import STRtree


REPOSITORY_ROOT = Path(__file__).resolve().parents[4]
DEFAULT_RAW = REPOSITORY_ROOT / "prototype" / "japan" / "raw-data" / "estat" / "od"
DEFAULT_BOUNDARY = None  # Resolve the selected World's ownership input at CLI entry.
DEFAULT_OUTPUT = REPOSITORY_ROOT / "map-creator" / "data" / "japan-prefectures"
DEFAULT_PREFECTURES = ("13", "14")
WORKER_VERSION = "estat-japan-ownership-boundary-v4"
PREFECTURE_NAMES = dict(zip(
    (f"{code:02d}" for code in range(1, 48)),
    (
        "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県",
        "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県", "新潟県", "富山県",
        "石川県", "福井県", "山梨県", "長野県", "岐阜県", "静岡県", "愛知県", "三重県",
        "滋賀県", "京都府", "大阪府", "兵庫県", "奈良県", "和歌山県", "鳥取県", "島根県",
        "岡山県", "広島県", "山口県", "徳島県", "香川県", "愛媛県", "高知県", "福岡県",
        "佐賀県", "長崎県", "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
    ),
    strict=True,
))


class Progress:
    def __init__(self, path: Path | None = None) -> None:
        self.path = path
        self.started = time.perf_counter()
        self.sequence = 0

    def emit(self, stage: str, status: str, **details: Any) -> None:
        self.sequence += 1
        event = {
            "event": "prefecture-progress",
            "sequence": self.sequence,
            "stage": stage,
            "status": status,
            "elapsedSeconds": round(time.perf_counter() - self.started, 3),
            "capturedAt": datetime.now(UTC).isoformat(),
            **details,
        }
        line = json.dumps(event, ensure_ascii=False, sort_keys=True)
        # Windows service consoles are commonly cp1252. Keep stdout valid JSON
        # there while retaining readable Unicode in the durable JSONL log.
        print(json.dumps(event, ensure_ascii=True, sort_keys=True), flush=True)
        if self.path:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            with self.path.open("a", encoding="utf-8", newline="\n") as output:
                output.write(line + "\n")


def number(value: object) -> int:
    """Convert published e-Stat values, including suppression markers, to ints."""
    if value is None:
        return 0
    text = str(value).strip().replace(",", "")
    if not text or text in {"*", "-", "X"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def mesh_center(code: str) -> tuple[float, float]:
    """Return longitude/latitude of a Japanese 500 m or 250 m mesh centre.

    8-digit third-level meshes are 1 km, 9-digit meshes are 500 m, and
    10-digit meshes are 250 m.  The downloaded source files use 9 and 10.
    """
    if not code.isdigit() or len(code) not in {9, 10}:
        raise ValueError(f"Unexpected Japanese mesh code: {code!r}")
    lat_seconds = int(code[:2]) * 2 * 60 * 60 // 3
    lon_seconds = (100 + int(code[2:4])) * 60 * 60
    lat_seconds += int(code[4]) * 5 * 60 + int(code[6]) * 30
    lon_seconds += int(code[5]) * 7.5 * 60 + int(code[7]) * 45
    if len(code) >= 9:
        quadrant = int(code[8])
        if quadrant not in {1, 2, 3, 4}:
            raise ValueError(f"Unexpected 500 m mesh quadrant: {code!r}")
        if quadrant in {3, 4}:
            lat_seconds += 15
        if quadrant in {2, 4}:
            lon_seconds += 22.5
    if len(code) == 9:
        lat_seconds += 7.5
        lon_seconds += 11.25
    if len(code) == 10:
        quadrant = int(code[9])
        if quadrant not in {1, 2, 3, 4}:
            raise ValueError(f"Unexpected 250 m mesh quadrant: {code!r}")
        if quadrant in {3, 4}:
            lat_seconds += 7.5
        if quadrant in {2, 4}:
            lon_seconds += 11.25
        lat_seconds += 3.75
        lon_seconds += 5.625
    return lon_seconds / 3600, lat_seconds / 3600


def polygon_parts(geometry: Any) -> list[Any]:
    if geometry.geom_type == "Polygon":
        return [geometry]
    if geometry.geom_type in {"MultiPolygon", "GeometryCollection"}:
        return [part for child in geometry.geoms for part in polygon_parts(child)]
    return []


class BoundaryOwnershipIndex:
    def __init__(self, selected: dict[str, Any], all_prefectures: dict[str, Any]) -> None:
        self.selected = selected
        self.selected_parts = []
        self.selected_codes = []
        self.selected_parts_by_code: dict[str, list[Any]] = {}
        self.selected_trees_by_code: dict[str, STRtree] = {}
        for pref_code, geometry in selected.items():
            parts = polygon_parts(geometry)
            self.selected_parts_by_code[pref_code] = parts
            self.selected_trees_by_code[pref_code] = STRtree(parts)
            self.selected_parts.extend(parts)
            self.selected_codes.extend([pref_code] * len(parts))
        self.selected_tree = STRtree(self.selected_parts)
        self.all_parts = []
        self.all_codes = []
        for pref_code, geometry in all_prefectures.items():
            parts = polygon_parts(geometry)
            self.all_parts.extend(parts)
            self.all_codes.extend([pref_code] * len(parts))
        self.all_tree = STRtree(self.all_parts)
        self._projection_cache: dict[str, tuple[Transformer, Transformer, Any]] = {}
        self._inset_cache: dict[tuple[str, int, float], Any] = {}

    def projected_boundary(self, pref_code: str) -> tuple[Transformer, Transformer, Any]:
        cached = self._projection_cache.get(pref_code)
        if cached:
            return cached
        boundary = self.selected[pref_code]
        center = boundary.centroid
        local_crs = (
            f"+proj=aeqd +lat_0={center.y:.12f} +lon_0={center.x:.12f} "
            "+datum=WGS84 +units=m +no_defs"
        )
        forward = Transformer.from_crs("EPSG:4326", local_crs, always_xy=True)
        inverse = Transformer.from_crs(local_crs, "EPSG:4326", always_xy=True)
        cached = (forward, inverse, transform(forward.transform, boundary))
        self._projection_cache[pref_code] = cached
        return cached


def load_prefecture_boundary(
    prefecture_codes: set[str],
    boundary_source: Path,
) -> tuple[dict[str, Any], dict[str, Any], BoundaryOwnershipIndex, dict[str, str]]:
    source = json.loads(boundary_source.read_text(encoding="utf-8"))
    if source.get("purpose") == "display-only" or "lods" in source:
        raise ValueError("Display LODs cannot determine demand ownership")
    features = [feature for feature in source["features"] if feature["properties"]["pref_code"] in prefecture_codes]
    found = {feature["properties"]["pref_code"] for feature in features}
    if found != prefecture_codes:
        raise ValueError(f"Missing prefecture boundaries: {sorted(prefecture_codes - found)}")
    all_prefectures = {
        feature["properties"]["pref_code"]: make_valid(shape(feature["geometry"]))
        for feature in source["features"]
    }
    selected = {code: all_prefectures[code] for code in prefecture_codes}
    names = {
        feature["properties"]["pref_code"]: str(
            feature["properties"].get("pref_name_ja")
            or feature["properties"].get("name")
            or feature["properties"]["pref_code"]
        )
        for feature in features
    }
    return (
        {"type": "FeatureCollection", "features": features},
        selected,
        BoundaryOwnershipIndex(selected, all_prefectures),
        names,
    )


def assign_prefecture(code: str, boundary_index: BoundaryOwnershipIndex) -> tuple[str, str] | None:
    """Assign a mesh against full World ownership, never a display LOD."""
    longitude, latitude = mesh_center(code)
    if len(code) not in (9, 10):
        raise ValueError(f"Unexpected Japanese mesh code: {code!r}")
    return resolve_cell_ownership(longitude, latitude, 1 if len(code) == 10 else 2, boundary_index)


def resolve_cell_ownership(longitude, latitude, mesh_scale, boundary_index):
    """One ownership rule for source extraction and final building placement.

    Prefer the centre's owner; for a coastal overhang use the greatest footprint
    overlap. Source prefecture labels and the current display zoom play no role.
    A cell wholly outside the World has no owner here, not an implicit fallback.
    """
    if mesh_scale not in (1, 2):
        raise ValueError('Expected a 250 m or 500 m mesh scale')
    center = Point(longitude, latitude)
    center_hits = boundary_index.all_tree.query(center, predicate="covered_by")
    if len(center_hits):
        pref_code = min(boundary_index.all_codes[int(index)] for index in center_hits)
        return pref_code, "center-inside-ownership-boundary"
    half_longitude, half_latitude = mesh_scale / 640, mesh_scale / 960
    cell = box(
        longitude - half_longitude,
        latitude - half_latitude,
        longitude + half_longitude,
        latitude + half_latitude,
    )
    overlaps: dict[str, float] = defaultdict(float)
    for index in boundary_index.all_tree.query(cell, predicate="intersects"):
        part_index = int(index)
        overlaps[boundary_index.all_codes[part_index]] += boundary_index.all_parts[part_index].intersection(cell).area
    if not overlaps:
        return None
    pref_code, overlap_area = min(overlaps.items(), key=lambda item: (-item[1], item[0]))
    if overlap_area <= 0:
        return None
    return pref_code, "maximum-ownership-boundary-overlap"


def relocate_into_boundary(
    longitude: float,
    latitude: float,
    pref_code: str,
    boundary_index: BoundaryOwnershipIndex,
) -> tuple[float, float, bool, float]:
    """Move a point to the closest covered location without changing its mass."""
    boundary = boundary_index.selected[pref_code]
    point = Point(longitude, latitude)
    if boundary.covers(point):
        return longitude, latitude, False, 0.0
    # Choose the location in the exact lon/lat geometry that is rendered. A
    # nearest point on a projected polygon is not reliable here: inverse-
    # projected straight chords do not reproduce GeoJSON's lon/lat edges.
    source_tree = boundary_index.selected_trees_by_code[pref_code]
    source_part_index = int(source_tree.nearest(point))
    source_part = boundary_index.selected_parts_by_code[pref_code][source_part_index]
    relocated = None
    for inset_degrees in (0.000002, 0.000005, 0.00001, 0.00002, 0.00005, 0.0001):
        cache_key = (pref_code, source_part_index, inset_degrees)
        target = boundary_index._inset_cache.get(cache_key)
        if target is None:
            target = source_part.buffer(-inset_degrees)
            boundary_index._inset_cache[cache_key] = target
        if target.is_empty:
            continue
        candidate = nearest_points(point, target)[1]
        serialized = Point(round(candidate.x, 7), round(candidate.y, 7))
        if boundary.covers(serialized):
            relocated = serialized
            break
    if relocated is None:
        candidate = source_part.representative_point()
        relocated = Point(round(candidate.x, 7), round(candidate.y, 7))
    if not boundary.covers(relocated):
        raise ValueError(f"Could not relocate ({longitude}, {latitude}) into rendered prefecture {pref_code}")
    forward, _, _ = boundary_index.projected_boundary(pref_code)
    original_x, original_y = forward.transform(longitude, latitude)
    relocated_x, relocated_y = forward.transform(relocated.x, relocated.y)
    distance_m = math.hypot(original_x - relocated_x, original_y - relocated_y)
    return relocated.x, relocated.y, True, round(distance_m, 3)


def rows_from_zip(path: Path, encoding: str = "cp932") -> Iterable[dict[str, str]]:
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist() if name.lower().endswith(".txt")]
        if len(names) != 1:
            raise ValueError(f"Expected one table in {path.name}, found {names}")
        with archive.open(names[0]) as handle:
            # The second human-readable title row is not part of the table.
            reader = csv.DictReader((line.decode(encoding) for line in handle))
            next(reader, None)
            yield from reader


def source_mesh_intersects(path: Path, selected_boundary: Any) -> bool:
    prefix = path.stem.rsplit("-M", 1)[-1]
    if len(prefix) != 4 or not prefix.isdigit():
        raise ValueError(f"Cannot read first-level mesh from {path.name}")
    min_latitude = int(prefix[:2]) * 2 / 3
    min_longitude = 100 + int(prefix[2:])
    return selected_boundary.intersects(
        box(min_longitude, min_latitude, min_longitude + 1, min_latitude + 2 / 3)
    )


def read_home_mesh(raw_root: Path, boundary_index: BoundaryOwnershipIndex, selected_boundary: Any) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    source_dir = raw_root / "2020-census-250m-origin-marginals"
    for source in sorted(source_dir.glob("*.zip")):
        if not source_mesh_intersects(source, selected_boundary):
            continue
        for row in rows_from_zip(source):
            # 0 is an unchanged released cell and 1 is a disclosure-merged
            # released cell.  2 is an input cell represented by its HTKSAKI.
            if row.get("HTKSYORI") not in {"0", "1"}:
                continue
            code = row["KEY_CODE"]
            commuters = number(row.get("T001109086"))
            if commuters <= 0:
                continue
            assignment = assign_prefecture(code, boundary_index)
            if assignment is None:
                continue
            pref_code, assignment_kind = assignment
            if pref_code not in boundary_index.selected:
                continue
            original_longitude, original_latitude = mesh_center(code)
            longitude, latitude, relocated, relocation_distance_m = relocate_into_boundary(
                original_longitude, original_latitude, pref_code, boundary_index
            )
            records.append({
                "id": f"home-250m-{code}", "meshCode": code,
                "longitude": longitude, "latitude": latitude,
                "commuters": commuters,
                "workers": number(row.get("T001109087")),
                "students": number(row.get("T001109088")),
                "railUsers": number(row.get("T001109090")),
                "prefCode": pref_code,
                "boundaryAssignment": assignment_kind,
                "relocated": relocated,
                "relocationDistanceM": relocation_distance_m,
                "originalLongitude": original_longitude,
                "originalLatitude": original_latitude,
            })
    return sorted(records, key=lambda record: record["meshCode"])


def read_job_mesh(raw_root: Path, boundary_index: BoundaryOwnershipIndex, selected_boundary: Any) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    source_dir = raw_root / "2021-economic-census-500m-destination-capacity"
    for source in sorted(source_dir.glob("*.zip")):
        if not source_mesh_intersects(source, selected_boundary):
            continue
        for row in rows_from_zip(source):
            code = row["KEY_CODE"]
            capacity = number(row.get("T001162022"))  # all-industry employees
            if capacity <= 0:
                continue
            assignment = assign_prefecture(code, boundary_index)
            if assignment is None:
                continue
            pref_code, assignment_kind = assignment
            if pref_code not in boundary_index.selected:
                continue
            original_longitude, original_latitude = mesh_center(code)
            longitude, latitude, relocated, relocation_distance_m = relocate_into_boundary(
                original_longitude, original_latitude, pref_code, boundary_index
            )
            records.append({
                "id": f"jobs-500m-{code}", "meshCode": code,
                "longitude": longitude, "latitude": latitude,
                "jobs": capacity,
                "establishments": number(row.get("T001162001")),
                "prefCode": pref_code,
                "boundaryAssignment": assignment_kind,
                "relocated": relocated,
                "relocationDistanceM": relocation_distance_m,
                "originalLongitude": original_longitude,
                "originalLatitude": original_latitude,
            })
    return sorted(records, key=lambda record: record["meshCode"])


def cluster_sites(
    records: list[dict[str, Any]],
    radius_m: float,
    boundary_index: BoundaryOwnershipIndex,
    id_prefix: str = "jp-prefecture-site",
) -> list[dict[str, Any]]:
    """Use the NEC maximal-radius seed rule, projected into metres for Japan."""
    if not records:
        return []
    prefecture_codes = {record["prefCode"] for record in records}
    if len(prefecture_codes) != 1:
        raise ValueError("Demand-site clustering must operate on exactly one prefecture")
    pref_code = next(iter(prefecture_codes))
    forward, _, _ = boundary_index.projected_boundary(pref_code)
    candidate_rows = []
    for record in records:
        x, y = forward.transform(record["longitude"], record["latitude"])
        candidate_rows.append({**record, "x": x, "y": y})
    candidates = sorted(candidate_rows, key=lambda row: (-row["commuters"], row["meshCode"]))
    seeds: list[dict[str, Any]] = []
    buckets: dict[tuple[int, int], list[dict[str, Any]]] = defaultdict(list)
    radius_squared = radius_m * radius_m
    for row in candidates:
        gx, gy = math.floor(row["x"] / radius_m), math.floor(row["y"] / radius_m)
        covered = any(
            (row["x"] - seed["x"]) ** 2 + (row["y"] - seed["y"]) ** 2 <= radius_squared
            for nx in range(gx - 1, gx + 2) for ny in range(gy - 1, gy + 2)
            for seed in buckets[(nx, ny)]
        )
        if not covered:
            seeds.append(row)
            buckets[(gx, gy)].append(row)
    members: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in candidate_rows:
        gx, gy = math.floor(row["x"] / radius_m), math.floor(row["y"] / radius_m)
        choices = [
            seed for nx in range(gx - 1, gx + 2) for ny in range(gy - 1, gy + 2)
            for seed in buckets[(nx, ny)]
            if (row["x"] - seed["x"]) ** 2 + (row["y"] - seed["y"]) ** 2 <= radius_squared
        ]
        if not choices:
            raise AssertionError(f"Seed set failed to cover {row['meshCode']}")
        seed = min(choices, key=lambda item: ((row["x"] - item["x"]) ** 2 + (row["y"] - item["y"]) ** 2, item["meshCode"]))
        members[seed["meshCode"]].append(row)
    sites = []
    for seed in seeds:
        grouped = members[seed["meshCode"]]
        total = sum(row["commuters"] for row in grouped)
        original_longitude = sum(row["longitude"] * row["commuters"] for row in grouped) / total
        original_latitude = sum(row["latitude"] * row["commuters"] for row in grouped) / total
        longitude, latitude, relocated, relocation_distance_m = relocate_into_boundary(
            original_longitude, original_latitude, pref_code, boundary_index
        )
        sites.append({
            "id": f"{id_prefix}-{seed['meshCode']}",
            "seedMeshCode": seed["meshCode"],
            "longitude": longitude,
            "latitude": latitude,
            "commuters": total,
            "workers": sum(row["workers"] for row in grouped),
            "students": sum(row["students"] for row in grouped),
            "railUsers": sum(row["railUsers"] for row in grouped),
            "sourceCellCount": len(grouped),
            "prefCode": seed["prefCode"],
            "relocated": relocated,
            "relocationDistanceM": relocation_distance_m,
            "originalLongitude": original_longitude,
            "originalLatitude": original_latitude,
        })
    return sorted(sites, key=lambda site: site["id"])


def relocation_summary(records: list[dict[str, Any]], mass_field: str) -> dict[str, Any]:
    relocated = [record for record in records if record["relocated"]]
    distances = sorted(float(record["relocationDistanceM"]) for record in relocated)
    return {
        "relocatedCount": len(relocated),
        "relocatedMass": sum(int(record[mass_field]) for record in relocated),
        "totalDistanceM": round(sum(distances), 3),
        "meanDistanceM": round(statistics.mean(distances), 3) if distances else 0,
        "maximumDistanceM": round(max(distances), 3) if distances else 0,
        "p95DistanceM": round(distances[math.ceil(len(distances) * 0.95) - 1], 3) if distances else 0,
    }


MUNICIPALITY = re.compile(r"^(\d{5})_")


def read_municipality_od(
    raw_root: Path,
    prefecture_codes: set[str],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Extract each selected origin's complete national destination vector."""
    from openpyxl import load_workbook

    workbook_path = raw_root / "2020-census-table-6-1-municipality-od-combined.xlsx"
    sheet = load_workbook(workbook_path, read_only=True, data_only=True).active
    header = next(sheet.iter_rows(min_row=8, max_row=8, values_only=True))
    destinations = []
    for index, value in enumerate(header):
        match = MUNICIPALITY.match(str(value or ""))
        # Codes ending in 00 are prefecture/designated-city aggregates.  Keep
        # the atomic municipalities and wards so totals are not double-counted.
        if match and not match.group(1).endswith("00"):
            destinations.append((index, match.group(1), str(value).split("_", 1)[1]))
    flows = []
    reconciliation = []
    for row in sheet.iter_rows(min_row=10, values_only=True):
        origin = str(row[3] or "")
        match = MUNICIPALITY.match(origin)
        if row[0] != "0_総数" or not match or match.group(1)[:2] not in prefecture_codes or match.group(1).endswith("00"):
            continue
        origin_code = match.group(1)
        origin_name = origin.split("_", 1)[1]
        emitted_cell_mass = 0
        for index, destination_code, destination_name in destinations:
            mass = number(row[index])
            if mass:
                emitted_cell_mass += mass
                flows.append({
                    "originMunicipalityCode": origin_code, "originName": origin_name,
                    "destinationMunicipalityCode": destination_code, "destinationName": destination_name,
                    "commutersAndStudents": mass,
                    "destinationKind": "unknown-residual" if destination_code == "99999" else "municipality",
                    "crossPrefecture": destination_code != "99999" and origin_code[:2] != destination_code[:2],
                })
        published_total = number(row[4])
        reconciliation.append({
            "originMunicipalityCode": origin_code,
            "publishedTotal": published_total,
            "emittedCellMass": emitted_cell_mass,
            "reconciliationDelta": emitted_cell_mass - published_total,
        })
    flows.sort(key=lambda flow: (flow["originMunicipalityCode"], flow["destinationMunicipalityCode"]))
    audit = {
        "originCount": len(reconciliation),
        "publishedOriginTotal": sum(row["publishedTotal"] for row in reconciliation),
        "acceptedAtomicMass": sum(row["emittedCellMass"] for row in reconciliation),
        "reconciliationDelta": sum(row["reconciliationDelta"] for row in reconciliation),
        "originReconciliation": reconciliation,
    }
    if sum(flow["commutersAndStudents"] for flow in flows) != audit["acceptedAtomicMass"]:
        raise AssertionError("Municipality O/D cells changed mass during extraction")
    return flows, audit


def point_collection(records: list[dict[str, Any]], properties: tuple[str, ...]) -> dict[str, Any]:
    return {"type": "FeatureCollection", "features": [
        {"type": "Feature", "properties": {key: record[key] for key in properties},
         "geometry": {"type": "Point", "coordinates": [round(record["longitude"], 7), round(record["latitude"], 7)]}}
        for record in records
    ]}


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8", newline="\n")
    os.replace(temporary, path)


def build_prefecture_evidence(
    *,
    prefecture_codes: list[str],
    raw_root: Path,
    boundary_source: Path,
    output: Path,
    site_radius_m: float,
    progress: Progress,
    queue_index: int | None = None,
    queue_total: int | None = None,
) -> dict[str, Any]:
    selected_codes = set(prefecture_codes)
    tile_ids = [f"JP_PREF_{code}" for code in prefecture_codes]
    common = {
        "workerVersion": WORKER_VERSION,
        "prefectureCodes": prefecture_codes,
        "prefectureNames": [PREFECTURE_NAMES[code] for code in prefecture_codes],
        "tileIds": tile_ids,
        "queueIndex": queue_index,
        "queueTotal": queue_total,
    }
    progress.emit("prefecture", "started", **common, output=str(output))
    boundary_geojson, boundaries, boundary_index, prefecture_names = load_prefecture_boundary(
        selected_codes,
        boundary_source,
    )
    selected_boundary = unary_union(list(boundaries.values()))
    progress.emit("boundary", "complete", **common, featureCount=len(boundary_geojson["features"]))

    homes = read_home_mesh(raw_root, boundary_index, selected_boundary)
    home_relocation = relocation_summary(homes, "commuters")
    progress.emit(
        "home-marginals",
        "complete",
        **common,
        cellCount=len(homes),
        mass=sum(row["commuters"] for row in homes),
        relocation=home_relocation,
    )
    jobs = read_job_mesh(raw_root, boundary_index, selected_boundary)
    job_relocation = relocation_summary(jobs, "jobs")
    progress.emit(
        "job-marginals",
        "complete",
        **common,
        cellCount=len(jobs),
        mass=sum(row["jobs"] for row in jobs),
        relocation=job_relocation,
    )
    sites = sorted(
        (
            site
            for pref_code in prefecture_codes
            for site in cluster_sites(
                [row for row in homes if row["prefCode"] == pref_code],
                site_radius_m,
                boundary_index,
                id_prefix=f"jp-pref-{pref_code}-site",
            )
        ),
        key=lambda site: site["id"],
    )
    site_relocation = relocation_summary(sites, "commuters")
    progress.emit(
        "demand-site-candidates",
        "complete",
        **common,
        candidateCount=len(sites),
        mass=sum(site["commuters"] for site in sites),
        relocation=site_relocation,
    )
    if sum(site["commuters"] for site in sites) != sum(row["commuters"] for row in homes):
        raise AssertionError("Demand-site clustering did not conserve home marginal mass")
    od_flows, od_reconciliation = read_municipality_od(raw_root, selected_codes)
    od_mass = sum(flow["commutersAndStudents"] for flow in od_flows)
    if od_mass != od_reconciliation["acceptedAtomicMass"]:
        raise AssertionError("Municipality O/D output did not conserve parsed atomic mass")
    progress.emit(
        "municipality-od",
        "complete",
        **common,
        flowCount=len(od_flows),
        mass=od_mass,
        publishedOriginTotal=od_reconciliation["publishedOriginTotal"],
        reconciliationDelta=od_reconciliation["reconciliationDelta"],
    )

    write_json(output / "world-boundary.geojson", boundary_geojson)
    write_json(output / "home-mesh-250m.geojson", point_collection(
        homes,
        (
            "id", "meshCode", "commuters", "workers", "students", "railUsers", "prefCode",
            "boundaryAssignment", "relocated", "relocationDistanceM", "originalLongitude", "originalLatitude",
        ),
    ))
    write_json(output / "job-mesh-500m.geojson", point_collection(
        jobs,
        (
            "id", "meshCode", "jobs", "establishments", "prefCode", "boundaryAssignment",
            "relocated", "relocationDistanceM", "originalLongitude", "originalLatitude",
        ),
    ))
    write_json(output / "demand-site-candidates.geojson", point_collection(
        sites,
        (
            "id", "seedMeshCode", "commuters", "workers", "students", "railUsers", "sourceCellCount",
            "prefCode", "relocated", "relocationDistanceM", "originalLongitude", "originalLatitude",
        ),
    ))
    write_json(output / "municipality-od.json", {
        "schemaVersion": 1,
        "workerVersion": WORKER_VERSION,
        "source": "e-Stat 2020 Census Table 6-1",
        "flows": od_flows,
        "reconciliation": od_reconciliation,
    })
    evidence = {
        "schemaVersion": 1,
        "workerVersion": WORKER_VERSION,
        "adapter": "estat-japan",
        "sourceVintage": "2020-2021",
        "partitions": tile_ids,
        "observations": [],
        "marginals": [
            *({"kind": "home", "location": row["id"], "mass": row["commuters"]} for row in homes),
            *({"kind": "job", "location": row["id"], "mass": row["jobs"]} for row in jobs),
        ],
        "odControls": [
            {
                "home": flow["originMunicipalityCode"],
                "work": flow["destinationMunicipalityCode"],
                "mass": flow["commutersAndStudents"],
            }
            for flow in od_flows
        ],
        "conservation": {"inputMass": od_mass, "acceptedMass": od_mass, "rejectedMass": 0},
    }
    write_json(output / "demand-evidence.json", evidence)
    cross_prefecture = sum(flow["commutersAndStudents"] for flow in od_flows if flow["crossPrefecture"])
    summary = {
        "schemaVersion": 1,
        "workerVersion": WORKER_VERSION,
        "status": "demand-evidence-complete",
        "generatedAt": datetime.now(UTC).isoformat(),
        "worldId": "JP_NATIONAL_OPEN_WORLD",
        "tileIds": tile_ids,
        "prefectures": prefecture_names,
        "ownershipBoundary": {
            "source": str(boundary_source),
            "sha256": sha256(boundary_source),
            "policy": "full approved World ownership geometry; independent of display LODs",
            "areaKm2": round(
                sum(boundary_index.projected_boundary(code)[2].area for code in prefecture_codes) / 1_000_000,
                2,
            ),
        },
        "sources": {
            "homeDemand": "2020 Census T001109, 250 m mesh; T001109086 residents aged 15+ who work or study",
            "jobCapacity": "2021 Economic Census T001162, 500 m mesh; T001162022 all-industry employees",
            "observedOd": "2020 Census Table 6-1, municipality residence-to-work/school-place matrix",
        },
        "candidateSites": {
            "algorithm": "deterministic maximal-radius seed set over official mesh marginals",
            "claim": "mesh-derived candidates; not building locations",
            "sourceMeshM": 250,
            "siteRadiusM": site_radius_m,
            "count": len(sites),
            "medianSourceCellsPerSite": statistics.median(site["sourceCellCount"] for site in sites) if sites else 0,
            "relocation": site_relocation,
        },
        "homeMesh": {
            "cellCount": len(homes),
            "commuters": sum(row["commuters"] for row in homes),
            "workers": sum(row["workers"] for row in homes),
            "students": sum(row["students"] for row in homes),
            "relocation": home_relocation,
        },
        "jobMesh": {
            "cellCount": len(jobs),
            "jobs": sum(row["jobs"] for row in jobs),
            "relocation": job_relocation,
        },
        "municipalityOd": {
            "flowCount": len(od_flows),
            "acceptedMass": od_mass,
            "crossPrefectureMass": cross_prefecture,
            "publishedOriginTotal": od_reconciliation["publishedOriginTotal"],
            "reconciliationDelta": od_reconciliation["reconciliationDelta"],
        },
        "boundaryAssignment": {
            "policy": "one authoritative World-boundary owner per source mesh; no fractional mass",
            "homeMaximumOverlap": sum(row["boundaryAssignment"] == "maximum-ownership-boundary-overlap" for row in homes),
            "jobMaximumOverlap": sum(row["boundaryAssignment"] == "maximum-ownership-boundary-overlap" for row in jobs),
        },
        "nextStage": "building-footprint allocation and map assets",
    }
    write_json(output / "report.json", summary)
    progress.emit(
        "prefecture",
        "complete",
        **common,
        homeCells=len(homes),
        jobCells=len(jobs),
        candidateSites=len(sites),
        odFlows=len(od_flows),
        odMass=od_mass,
        relocation={"home": home_relocation, "job": job_relocation, "candidateSites": site_relocation},
        artifact=str(output / "demand-evidence.json"),
    )
    return summary


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prefecture", action="append", dest="prefectures", help="two-digit prefecture code; repeatable")
    parser.add_argument("--raw-root", type=Path, default=DEFAULT_RAW)
    parser.add_argument("--world-root", type=Path, default=REPOSITORY_ROOT / "worlds" / "japan")
    parser.add_argument(
        "--boundary-source",
        type=Path,
        default=DEFAULT_BOUNDARY,
        help="full approved World ownership geometry; never a display LOD",
    )
    parser.add_argument("--output", type=Path)
    parser.add_argument("--progress-jsonl", type=Path)
    parser.add_argument("--queue-index", type=int)
    parser.add_argument("--queue-total", type=int)
    parser.add_argument("--site-radius-m", type=float, default=350, help="NEC-style Voronoi seed coverage radius in metres (default: 350)")
    args = parser.parse_args(argv)
    if args.boundary_source is None:
        from ..geography import ownership_boundary
        args.boundary_source = ownership_boundary(args.world_root)
    if args.site_radius_m < 250:
        parser.error("site radius must be at least 250 m: smaller values out-resolve the source mesh")
    prefecture_codes = args.prefectures or list(DEFAULT_PREFECTURES)
    if any(len(code) != 2 or not code.isdigit() or not 1 <= int(code) <= 47 for code in prefecture_codes):
        parser.error("--prefecture must be a two-digit code from 01 through 47")
    if len(set(prefecture_codes)) != len(prefecture_codes):
        parser.error("--prefecture values must be unique")
    output = args.output or DEFAULT_OUTPUT / (f"JP_PREF_{prefecture_codes[0]}" if len(prefecture_codes) == 1 else "_".join(prefecture_codes))
    progress_path = args.progress_jsonl or output / "progress.jsonl"
    if progress_path.exists():
        progress_path.unlink()
    progress = Progress(progress_path)
    try:
        summary = build_prefecture_evidence(
            prefecture_codes=prefecture_codes,
            raw_root=args.raw_root,
            boundary_source=args.boundary_source,
            output=output,
            site_radius_m=args.site_radius_m,
            progress=progress,
            queue_index=args.queue_index,
            queue_total=args.queue_total,
        )
    except Exception as error:
        progress.emit(
            "prefecture",
            "failed",
            prefectureCodes=prefecture_codes,
            errorType=type(error).__name__,
            error=str(error),
        )
        raise
    print(json.dumps(summary, ensure_ascii=True, sort_keys=True), flush=True)


if __name__ == "__main__":
    main()
